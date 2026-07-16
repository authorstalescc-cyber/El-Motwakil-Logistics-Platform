# -*- coding: utf-8 -*-
"""
تصدير بيانات النظام القديم (Excel) إلى ملف JSON نظيف
يستهلكه سكريبت التثبيت PHP فيدخلها إلى MySQL بحسابات مطابقة تماماً.
"""
import json, os, sys, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
OLD = os.environ.get("OLD_XLSX", os.path.join(HERE, "..", "..", "old-system-data.xlsx"))
OUT = os.path.join(HERE, "old_data.json")

try:
    import openpyxl
except ImportError:
    print("openpyxl غير مثبت. ثبّته: pip3 install openpyxl", file=sys.stderr); sys.exit(1)

if not os.path.isfile(OLD):
    print(f"ملف النظام القديم غير موجود: {OLD}", file=sys.stderr); sys.exit(1)

def clean(v):
    if v is None: return None
    if isinstance(v, str):
        v = v.strip()
        if v == "" or v.startswith("#"): return None
        return v
    return v

def num(v):
    v = clean(v)
    if v is None: return 0.0
    try: return float(v)
    except (ValueError, TypeError): return 0.0

def dt(v):
    v = clean(v)
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return None

wb = openpyxl.load_workbook(OLD, data_only=True, read_only=True)

def rows_of(sheet, maxrow):
    if sheet not in wb.sheetnames:
        # محاولة مطابقة مرنة (فروق المسافات/الياء)
        for s in wb.sheetnames:
            if s.strip().replace("ى","ي") == sheet.strip().replace("ى","ي"):
                sheet = s; break
        else:
            print(f"تحذير: الصفحة '{sheet}' غير موجودة", file=sys.stderr); return []
    ws = wb[sheet]
    return list(ws.iter_rows(min_row=2, max_row=maxrow, values_only=True))

def cell(row, idx):
    return row[idx-1] if len(row) >= idx else None

# ---------------- البوالص ----------------
bawalis = []
for row in rows_of("بوالص", 4041):
    date = dt(cell(row,2)); bill = clean(cell(row,3))
    if date is None and bill is None:
        continue
    other = num(cell(row,29)) + num(cell(row,30)) + num(cell(row,31))
    bawalis.append({
        "declaration_owner": clean(cell(row,1)),
        "bill_date": date,
        "bill_number": str(bill) if bill is not None else None,
        "classification": clean(cell(row,4)),
        "company": clean(cell(row,5)),
        "destination": clean(cell(row,6)),
        "container_no": clean(cell(row,7)),
        "car_no": clean(cell(row,8)),
        "driver": clean(cell(row,9)),
        "car_owner": clean(cell(row,10)),
        "freight": num(cell(row,11)),
        "custody": num(cell(row,12)),
        "transfer": num(cell(row,13)),
        "discount": num(cell(row,14)),
        "overnight_collection": num(cell(row,15)),
        "payment": num(cell(row,17)),
        "driver_share": num(cell(row,19)),
        "assistant_share": num(cell(row,20)),
        "overnight_expense": num(cell(row,21)),
        "expense": num(cell(row,22)),
        "cards": num(cell(row,23)),
        "tahia_masr": num(cell(row,24)),
        "thabtan": num(cell(row,25)),
        "z_kawtsh": num(cell(row,26)),
        "permits": num(cell(row,27)),
        "kawtsh": num(cell(row,28)),
        "other_expense": other,
        "gas": num(cell(row,32)),
    })

# ---------------- التخليص ----------------
takhlis = []
for row in rows_of("تخليص", 338):
    date = dt(cell(row,2)); name = clean(cell(row,3)); cnum = clean(cell(row,4))
    fees = num(cell(row,5)); exp = num(cell(row,6)); pay = num(cell(row,9))
    if date is None and name is None and cnum is None and fees==0 and exp==0 and pay==0:
        continue
    takhlis.append({
        "cert_date": date, "cert_name": name,
        "cert_number": str(cnum) if cnum is not None else None,
        "fees": fees, "expenses": exp, "payments": pay,
    })

# ---------------- المولدات ----------------
def read_gen(sheet, maxrow, source):
    out = []
    for row in rows_of(sheet, maxrow):
        date = dt(cell(row,2)); comp = clean(cell(row,3))
        if date is None and comp is None:
            continue
        out.append({
            "source": source, "gen_date": date, "company": comp,
            "address": clean(cell(row,4)), "permit_no": clean(cell(row,5)),
            "gen_count": int(num(cell(row,6))), "freight": num(cell(row,7)),
            "payments": num(cell(row,8)),
        })
    return out
generators = read_gen("مولدات ايجى كول ", 340, "egycool") + read_gen("مولدات الشجيع", 25, "shujaie")

# ---------------- الأرصدة ----------------
balances = []
if "اجمالي" in wb.sheetnames:
    ag = wb["اجمالي"]
    for row in ag.iter_rows(min_row=2, max_row=60, values_only=True):
        cname = clean(cell(row,1)); cbal = num(cell(row,2))
        sname = clean(cell(row,3)); sbal = num(cell(row,4))
        if cname: balances.append({"party_type":"client","name":str(cname),"balance":cbal})
        if sname: balances.append({"party_type":"supplier","name":str(sname),"balance":sbal})

wb.close()

# ---------------- القوائم المنسدلة (قيم مميّزة) ----------------
def distinct(items, key):
    seen = []
    for it in items:
        v = it.get(key)
        if v and isinstance(v, str) and v.strip() and v not in seen:
            seen.append(v.strip())
    return sorted(seen)

lookups = {
    "classification": distinct(bawalis, "classification"),
    "company": sorted(set(distinct(bawalis,"company") + distinct(generators,"company"))),
    "destination": distinct(bawalis, "destination"),
    "driver": distinct(bawalis, "driver"),
    "car_owner": distinct(bawalis, "car_owner"),
    "client": sorted(set([b["name"] for b in balances if b["party_type"]=="client"])),
}

data = {"bawalis": bawalis, "takhlis": takhlis, "generators": generators,
        "balances": balances, "lookups": lookups}

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(data, f, ensure_ascii=False)

print(f"تم التصدير → {OUT}")
print(f"  بوالص={len(bawalis)}  تخليص={len(takhlis)}  مولدات={len(generators)}  أرصدة={len(balances)}")
print(f"  قوائم: " + "  ".join(f"{k}={len(v)}" for k,v in lookups.items()))
