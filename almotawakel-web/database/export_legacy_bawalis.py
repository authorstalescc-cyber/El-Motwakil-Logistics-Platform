# -*- coding: utf-8 -*-
"""
استيراد البوالص من النظام القديم (Excel) وتنظيفها → JSON نظيف.

  * لا يمسّ أي حساب في البرنامج الحالي — يُصدَّر إلى جدول أرشيف مستقل (legacy_bawalis)
    عبر سكربت PHP  bin/import_legacy_bawalis.php.
  * يوحّد المسميات:
      - "خ" في آخر «صاحب البيان» = علامة تكرار تُزال ("احمد عبدو خ" → "احمد عبدو").
      - "ق" في آخر «صاحب البيان» كتوكن منفصل تُزال فقط لو الاسم الأصلي (بدونها) موجود.
      - توحيد التهجئة (ي/ى، الهمزات، ة/ه، التشكيل، المسافات) بنفس منطق normkey في PHP،
        مع اختيار أكثر صيغة شائعة كاسم قانوني لكل مجموعة (دمج المكرّرات).
      - "؟؟" وما شابه → فارغ.
  * يستورد كل حقول البوليصة (نولون/عهدة/مصروفات/باقي/…) بأمانة للأرشيف.

التشغيل:  python3 database/export_legacy_bawalis.py
البيئة:   OLD_XLSX لتحديد مسار ملف Excel (الافتراضي: ../old-system-data.xlsx)
"""
import json, os, sys, datetime, re
from collections import Counter, defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
OLD  = os.environ.get("OLD_XLSX", os.path.join(HERE, "..", "..", "old-system-data.xlsx"))
OUT  = os.path.join(HERE, "legacy_bawalis.json")

try:
    import openpyxl
except ImportError:
    print("openpyxl غير مثبت. ثبّته: pip3 install openpyxl", file=sys.stderr); sys.exit(1)
if not os.path.isfile(OLD):
    print(f"ملف النظام القديم غير موجود: {OLD}", file=sys.stderr); sys.exit(1)

# ───────────────────────── أدوات التنظيف ─────────────────────────
def clean(v):
    """قصّ + تحويل الفراغ/الرموز غير المفيدة إلى None"""
    if v is None: return None
    if isinstance(v, str):
        v = v.strip()
        if v == "" or v in ("؟؟", "؟", "??", "#", "-", "—"): return None
        if v.startswith("#"): return None
        return v
    return v

def num(v):
    v = clean(v)
    if v is None: return 0.0
    try: return round(float(v), 2)
    except (ValueError, TypeError): return 0.0

_year_fixes = []   # تواريخ سنة 1900 (خطأ إدخال إكسل) صُحّحت إلى 2026

def dt(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        # سنة 1900 = خطأ إدخال في الإكسل (خانة السنة فاضية/قيمة مسحوبة) —
        # كل بيانات الأرشيف 2026، واليوم/الشهر متسقان مع تسلسل الإدخال.
        if v.year == 1900:
            _year_fixes.append(v)
            v = v.replace(year=2026)
        return v.strftime("%Y-%m-%d")
    return None

# نفس منطق App\Models\Maintenance::normkey (لتوحيد التهجئة عند المقارنة فقط)
_HARAKAT = re.compile(r'[ـً-ْٰ]')     # تطويل + تشكيل + ألف خنجرية
_ALEF    = re.compile(r'[أإآٱ]')      # أ إ آ ٱ
def normkey(s):
    if s is None: return ''
    s = str(s).strip()
    if s == '': return ''
    s = _HARAKAT.sub('', s)
    s = _ALEF.sub('ا', s)
    s = s.translate(str.maketrans({'ى':'ي', 'ة':'ه', 'ؤ':'و', 'ئ':'ي'}))
    s = re.sub(r'\s+', ' ', s).strip()
    return s.lower()

def tidy(s):
    """تنظيف سطحي للعرض: قصّ + توحيد المسافات (بدون تغيير الحروف)"""
    if s is None: return None
    s = re.sub(r'\s+', ' ', str(s)).strip()
    return s or None

# ───────────────────────── قراءة الشيت ─────────────────────────
wb = openpyxl.load_workbook(OLD, data_only=True, read_only=True)
if "بوالص" not in wb.sheetnames:
    print("الصفحة 'بوالص' غير موجودة", file=sys.stderr); sys.exit(1)
ws = wb["بوالص"]
raw = list(ws.iter_rows(min_row=2, values_only=True))   # حتى آخر صف فعلي (بلا حد ثابت)
wb.close()

def cell(row, i):     # 1-based
    return row[i-1] if len(row) >= i else None

# صفوف صالحة: تاريخ أو رقم بوليصة موجود
records = []
for ri, row in enumerate(raw, start=2):
    date = dt(cell(row, 2)); bill = clean(cell(row, 3))
    if date is None and bill is None:
        continue
    records.append((ri, row, date, bill))

# ───────────── معالجة «صاحب البيان»: إزالة علامتَي خ/ق ─────────────
def strip_owner_marker(v, base_keys):
    """يزيل توكن 'خ' النهائي دائماً، و'ق' النهائي المنفصل لو الأساس موجود."""
    if v is None: return None
    v = tidy(v)
    if v is None: return None
    toks = v.split(' ')
    if len(toks) >= 2 and toks[-1] == 'خ':               # "... خ" منفصلة
        toks = toks[:-1]
    elif len(toks) >= 1 and toks[-1].endswith('خ') and toks[-1][:-1].isdigit():  # "1خ" ملتصقة برقم
        toks[-1] = toks[-1][:-1]
    v2 = ' '.join(t for t in toks if t).strip()
    # علامة "ق" المنفصلة: تُزال فقط لو الأساس (بدونها) اسم معروف
    toks2 = v2.split(' ')
    if len(toks2) >= 2 and toks2[-1] == 'ق':
        base = ' '.join(toks2[:-1]).strip()
        if normkey(base) in base_keys:
            v2 = base
    return v2 or None

# أساس أسماء أصحاب البيان (بعد إزالة خ) — لاختبار وجود أساس ق
owner_base_keys = set()
for _, row, _, _ in records:
    v = tidy(cell(row, 1))
    if v is None: continue
    toks = v.split(' ')
    if len(toks) >= 2 and toks[-1] == 'خ': toks = toks[:-1]
    v1 = ' '.join(toks).strip()
    if v1 and not (v1.split(' ')[-1] == 'ق'):
        owner_base_keys.add(normkey(v1))

# ───────────── توحيد المسميات: خريطة قانونية لكل عمود ─────────────
# لكل عمود اسمي: جمّع القيم بمفتاح normkey، واختر أكثر صيغة شائعة كاسم قانوني.
NAME_COLS = {
    'declaration_owner': 1,
    'classification':    4,
    'company':           5,
    'destination':       6,
    'driver':            9,
    'car_owner':         10,
}

# القيمة المُنظّفة الأولية لكل خلية اسمية (قبل التوحيد)
def pre_value(col_key, row):
    idx = NAME_COLS[col_key]
    v = clean(cell(row, idx))
    if v is None: return None
    if col_key == 'declaration_owner':
        return strip_owner_marker(v, owner_base_keys)
    return tidy(v)

# بناء الخرائط القانونية
canon = {}          # col_key -> { normkey -> canonical surface }
merge_report = {}   # col_key -> list of (canonical, [variants])
for col_key in NAME_COLS:
    groups = defaultdict(Counter)   # normkey -> Counter(surface -> count)
    for _, row, _, _ in records:
        v = pre_value(col_key, row)
        if v is None: continue
        groups[normkey(v)][v] += 1
    cmap = {}; report = []
    for k, counter in groups.items():
        # القانوني = الأكثر تكراراً، ثم الأقصر، ثم أبجدياً
        best = sorted(counter.items(), key=lambda kv: (-kv[1], len(kv[0]), kv[0]))[0][0]
        cmap[k] = best
        if len(counter) > 1:
            report.append((best, sorted(counter.keys())))
    canon[col_key] = cmap
    merge_report[col_key] = report

def canonical(col_key, row):
    v = pre_value(col_key, row)
    if v is None: return None
    return canon[col_key].get(normkey(v), v)

# ───────────── بناء السجلات النهائية (كل الحقول) ─────────────
out = []
for ri, row, date, bill in records:
    out.append({
        "src_row":            ri,
        "declaration_owner":  canonical('declaration_owner', row),
        "bill_date":          date,
        "bill_number":        str(bill) if bill is not None else None,
        "classification":     canonical('classification', row),
        "company":            canonical('company', row),
        "destination":        canonical('destination', row),
        "container_no":       tidy(clean(cell(row, 7))),
        "car_no":             tidy(clean(cell(row, 8))),
        "driver":             canonical('driver', row),
        "car_owner":          canonical('car_owner', row),
        "freight":            num(cell(row, 11)),   # نولون
        "custody":            num(cell(row, 12)),   # عهدة
        "transfer":           num(cell(row, 13)),   # تحويلة
        "discount":           num(cell(row, 14)),   # خصم
        "overnight_collection": num(cell(row, 15)), # بياتة (تحصيل)
        "remaining":          num(cell(row, 16)),   # باقى
        "payment":            num(cell(row, 17)),   # دفعه
        "total":              num(cell(row, 18)),   # اجمالى
        "driver_share":       num(cell(row, 19)),   # دور سواق
        "assistant_share":    num(cell(row, 20)),   # دور تباع
        "overnight_expense":  num(cell(row, 21)),   # بياتة (مصروف)
        "expense":            num(cell(row, 22)),   # مصروف
        "cards":              num(cell(row, 23)),   # كرتات
        "tahia_masr":         num(cell(row, 24)),   # تحيا مصر
        "thabtan":            num(cell(row, 25)),   # ثبتان
        "z_kawtsh":           num(cell(row, 26)),   # ظ كاوتش
        "permits_fee":        num(cell(row, 27)),   # تصاريح
        "kawtsh":             num(cell(row, 28)),   # كاوتش
        "other_expense":      round(num(cell(row, 29)) + num(cell(row, 30)) + num(cell(row, 31)), 2),  # م اخرى ×3
        "gas":                num(cell(row, 32)),   # جاز
        "total_expense":      num(cell(row, 33)),   # اجمال المصروف
        "custody_remaining":  num(cell(row, 34)),   # الباقى من العهدة
        "car_net":            num(cell(row, 35)),   # صافى السيارة
    })

with open(OUT, "w", encoding="utf-8") as f:
    json.dump({"bawalis": out}, f, ensure_ascii=False)

# ───────────────────────── التقرير ─────────────────────────
def col_stats(col_key):
    """الأصل (قبل أي تنظيف) → النهائي (بعد خ/ق + توحيد التهجئة)"""
    idx = NAME_COLS[col_key]
    orig_vals = set(); final_vals = set()
    for _, row, _, _ in records:
        ov = clean(cell(row, idx))
        if ov: orig_vals.add(tidy(ov))
        c = canonical(col_key, row)
        if c: final_vals.add(c)
    return len(orig_vals), len(final_vals)

print(f"✓ تم التصدير → {OUT}")
print(f"  عدد البوالص: {len(out)}")
print(f"  Σ النولون = {sum(r['freight'] for r in out):,.0f}   Σ العهدة = {sum(r['custody'] for r in out):,.0f}")
print("\n  توحيد المسميات (قيم مميّزة قبل → بعد):")
LABELS = {'declaration_owner':'صاحب البيان','classification':'التصنيف','company':'الشركة',
          'destination':'الجهة','driver':'السائق','car_owner':'صاحب السيارة'}
for col_key in NAME_COLS:
    b, a = col_stats(col_key)
    print(f"    {LABELS[col_key]:14} {b:4} → {a:4}   (دُمج {b-a})")

print("\n  إزالة علامة خ/ق من «صاحب البيان» (الأصل → بعد الدمج):")
kh_report = {}
for _, row, _, _ in records:
    ov = tidy(clean(cell(row, 1)))
    if ov is None: continue
    nv = canonical('declaration_owner', row)
    if nv and nv != ov and (ov.split(' ')[-1] in ('خ', 'ق') or ov.endswith('خ')):
        kh_report[ov] = nv
for ov, nv in sorted(kh_report.items()):
    print(f"    «{ov}»  →  «{nv}»")
