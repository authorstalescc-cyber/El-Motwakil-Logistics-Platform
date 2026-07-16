# -*- coding: utf-8 -*-
"""
مزامنة القوائم: يحدّث (الشركات / الجهات / ملاك السيارات) في ملف التقفيل اليومي
من الملف الأساسي — بدون أن يلمس أي بيانات أدخلتها (الرحلات/الإيصالات/سجل اليومية).

الاستخدام:
    1) اقفل ملف (المتوكل - تقفيل اليومية.xlsx) في Excel.
    2) شغّل:  python3 sync_lists.py
"""
import os
from openpyxl import load_workbook

DIR = "/Users/karimmokhtar/المتوكل"
MAIN = os.path.join(DIR, "المتوكل - نظام إدارة النقل والتخليص.xlsx")
DAILY = os.path.join(DIR, "المتوكل - تقفيل اليومية.xlsx")

SET_SHEET = "الإعدادات"
START, END = 4, 203
# (عمود المصدر في الأساسي) -> (عمود الوجهة في التقفيل)
MAP = {"E": "H",   # الشركات
       "F": "I",   # الجهات
       "H": "J"}   # ملاك السيارات

def main():
    mwb = load_workbook(MAIN, data_only=True, read_only=True)
    mset = mwb[SET_SHEET]
    src = {}
    for scol in MAP:
        src[scol] = [mset[f"{scol}{r}"].value for r in range(START, END + 1)]
    mwb.close()

    dwb = load_workbook(DAILY)
    dset = dwb[SET_SHEET]
    n = {}
    for scol, dcol in MAP.items():
        vals = src[scol]
        cnt = 0
        for i, r in enumerate(range(START, END + 1)):
            v = vals[i]
            dset[f"{dcol}{r}"] = (v if v not in (None, "") else None)
            if v not in (None, ""):
                cnt += 1
        n[dcol] = cnt
    dwb.save(DAILY)
    print("تمت المزامنة ✓  الشركات=%d  الجهات=%d  الملاك=%d" % (n["H"], n["I"], n["J"]))

if __name__ == "__main__":
    main()
