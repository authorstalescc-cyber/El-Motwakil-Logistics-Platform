# -*- coding: utf-8 -*-
"""
نظام المتوكل لإدارة النقل والتخليص الجمركي
مولّد ملف الإكسيل الاحترافي — openpyxl
"""
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule, DataBarRule
from openpyxl.chart import BarChart, PieChart, LineChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.properties import PageSetupProperties

# ============================ الألوان والهوية ============================
NAVY       = "FF15355B"   # الكحلي الأساسي
NAVY_DARK  = "FF0E2440"
NAVY_MID   = "FF1E4B7A"
GOLD       = "FFC9A227"   # الذهبي (لون الهوية)
GOLD_LIGHT = "FFF3E9C6"
GREEN      = "FF1E6E3C"   # الإيرادات
GREEN_LT   = "FFE3F1E7"
RED        = "FF9B2C2C"   # المصروفات
RED_LT     = "FFF6E4E4"
TEAL       = "FF14746F"
PURPLE     = "FF5B3E7A"
WHITE      = "FFFFFFFF"
INK        = "FF1A2433"
GREY_TXT   = "FF5A6472"
INPUT_FILL = "FFFFFDF3"   # خلية إدخال (أصفر فاتح جداً)
CALC_FILL  = "FFEDF2F8"   # خلية محسوبة (أزرق فاتح جداً)
ZEBRA      = "FFF7F9FB"
GRID       = "FFD5DBE2"
BANNER_BG  = NAVY
SECTION_BG = "FFEAF0F6"

FONT_NAME  = "Arial"

# صيغ الأرقام
MONEY = '#,##0.00 "ج.م"'
MONEY0 = '#,##0 "ج.م"'
PCT   = '0.0%'
DATEF = 'DD/MM/YYYY'
INTF  = '#,##0'

DATA_ROWS = 2200         # عدد صفوف الإدخال المجهّزة بكل صفحة (تكفي بيانات النظام القديم + نمو)
HDR_ROW   = 3            # صف رؤوس الأعمدة
DATA_START = 4           # أول صف بيانات
DATA_END  = DATA_START + DATA_ROWS - 1   # آخر صف بيانات

thin = Side(style="thin", color=GRID)
med  = Side(style="medium", color=NAVY)
box_thin = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ============================ دوال مساعدة ============================
def sref(sheet):
    """اسم صفحة بين علامتي اقتباس للاستخدام في المعادلات"""
    return "'" + sheet + "'"

def col_range(sheet, col_letter):
    return f"{sref(sheet)}!${col_letter}${DATA_START}:${col_letter}${DATA_END}"

def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def banner(ws, ncols, title, subtitle=""):
    """شريط عنوان علوي احترافي عبر صفين"""
    last = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    c = ws["A1"]
    c.value = "شركة المتوكل  —  " + title
    c.font = Font(name=FONT_NAME, size=20, bold=True, color=GOLD)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.fill = PatternFill("solid", fgColor=BANNER_BG)
    s = ws["A2"]
    s.value = subtitle if subtitle else "لخدمات النقل والتخليص الجمركي"
    s.font = Font(name=FONT_NAME, size=11, bold=False, color=WHITE)
    s.alignment = Alignment(horizontal="center", vertical="center")
    s.fill = PatternFill("solid", fgColor=NAVY_MID)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 20
    # تلوين خلايا الشريط المدمجة
    for r in (1, 2):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=r, column=col)
            cell.fill = PatternFill("solid", fgColor=(BANNER_BG if r == 1 else NAVY_MID))

def style_headers(ws, headers, group_colors):
    """كتابة رؤوس الأعمدة الملوّنة حسب المجموعة"""
    for i, (h, gc) in enumerate(zip(headers, group_colors), start=1):
        cell = ws.cell(row=HDR_ROW, column=i, value=h)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = PatternFill("solid", fgColor=gc)
        cell.border = Border(left=Side(style="thin", color=WHITE),
                             right=Side(style="thin", color=WHITE),
                             top=thin, bottom=Side(style="medium", color=GOLD))
    ws.row_dimensions[HDR_ROW].height = 40

def style_data_cell(cell, kind, numfmt, zebra=False):
    fill = INPUT_FILL if kind == "input" else CALC_FILL
    if zebra and kind == "input":
        fill = ZEBRA
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.border = box_thin
    cell.font = Font(name=FONT_NAME, size=10,
                     color=INK, bold=False,
                     italic=(kind == "formula"))
    cell.alignment = Alignment(horizontal="center", vertical="center")
    if numfmt:
        cell.number_format = numfmt

def finalize_entry_sheet(ws, ncols):
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = f"A{DATA_START}"
    last = get_column_letter(ncols)
    ws.auto_filter.ref = f"A{HDR_ROW}:{last}{DATA_END}"
    ws.sheet_view.showGridLines = False
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


# ملخص أسفل كل صفحة (صف الإجماليات)
def totals_row(ws, ncols, sum_cols, label_col=2, label="الإجماليات"):
    r = DATA_END + 1
    ws.row_dimensions[r].height = 22
    for col in range(1, ncols + 1):
        cell = ws.cell(row=r, column=col)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color=GOLD)
        cell.border = Border(top=Side(style="double", color=GOLD))
        cell.alignment = Alignment(horizontal="center", vertical="center")
    lc = ws.cell(row=r, column=label_col)
    lc.value = label
    for (col, numfmt) in sum_cols:
        L = get_column_letter(col)
        c = ws.cell(row=r, column=col)
        c.value = f"=SUM(${L}${DATA_START}:${L}${DATA_END})"
        c.number_format = numfmt
    return r


# ==========================================================================
#                        (1) صفحة الإعدادات والقوائم
# ==========================================================================
S = "الإعدادات"
ws_set = wb.active
ws_set.title = S
ws_set.sheet_view.rightToLeft = True
ws_set.sheet_view.showGridLines = False

banner(ws_set, 8, "الإعدادات والقوائم المنسدلة", "تُغذّي هذه الصفحة كل القوائم المنسدلة في النظام — أضف أي عنصر جديد هنا")

# إعدادات عامة (يسار)
gen = [
    ("اسم الشركة", "شركة المتوكل"),
    ("نسبة القيمة المضافة", 0.14),
    ("سنة التقارير", 2026),
    ("العملة", "ج.م"),
]
ws_set["A3"] = "الإعداد"; ws_set["B3"] = "القيمة"
for cell in (ws_set["A3"], ws_set["B3"]):
    cell.font = Font(name=FONT_NAME, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = box_thin
for i, (k, v) in enumerate(gen):
    r = 4 + i
    a = ws_set.cell(row=r, column=1, value=k)
    b = ws_set.cell(row=r, column=2, value=v)
    a.font = Font(name=FONT_NAME, bold=True, color=INK)
    a.fill = PatternFill("solid", fgColor=SECTION_BG)
    b.font = Font(name=FONT_NAME, bold=True, color=GREEN)
    b.fill = PatternFill("solid", fgColor=INPUT_FILL)
    for c in (a, b):
        c.border = box_thin
        c.alignment = Alignment(horizontal="center", vertical="center")
ws_set["B5"].number_format = PCT
ws_set["B6"].number_format = "0"

# القوائم المنسدلة (تبدأ من العمود D)
LISTS = {
    "التصنيفات": ("D", ["مستورد", "مصدّر", "ترانزيت", "بضائع عامة", "مواد خام", "معدات", "أخرى"]),
    "الشركات":   ("E", ["", "", ""]),
    "الجهات":    ("F", ["الإسكندرية", "الدخيلة", "العين السخنة", "دمياط", "القاهرة", "بورسعيد", "السويس", "6 أكتوبر", "العاشر من رمضان"]),
    "السائقين":  ("G", ["", "", ""]),
    "أصحاب السيارات": ("H", ["", "", ""]),
    "العملاء":   ("I", ["", "", ""]),
}
LIST_MAX = 200   # أقصى عدد عناصر بكل قائمة
list_col_map = {}
for name, (col, seed) in LISTS.items():
    list_col_map[name] = col
    h = ws_set[f"{col}3"]
    h.value = name
    h.font = Font(name=FONT_NAME, bold=True, color=WHITE)
    h.fill = PatternFill("solid", fgColor=GOLD)
    h.alignment = Alignment(horizontal="center", vertical="center")
    h.border = box_thin
    ws_set.column_dimensions[col].width = 20
    for j in range(LIST_MAX):
        cell = ws_set.cell(row=4 + j, column=ord(col) - 64)
        if j < len(seed):
            cell.value = seed[j] if seed[j] else None
        cell.fill = PatternFill("solid", fgColor=INPUT_FILL)
        cell.border = Border(left=thin, right=thin, bottom=thin, top=thin)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = Font(name=FONT_NAME, size=10, color=INK)

set_widths(ws_set, [24, 16, 3, 20, 22, 20, 22, 20])
ws_set.column_dimensions["C"].width = 3

# الأسماء المعرّفة (Named ranges)
def dn(name, ref):
    wb.defined_names.add(DefinedName(name, attr_text=ref))

dn("نسبة_القيمة_المضافة", f"{sref(S)}!$B$5")
dn("سنة_التقارير", f"{sref(S)}!$B$6")
for name, col in list_col_map.items():
    dn("قائمة_" + name.replace(" ", "_"), f"{sref(S)}!${col}$4:${col}${3+LIST_MAX}")


def add_list_validation(ws, col_letter, list_name):
    dv = DataValidation(type="list",
                        formula1=f"={sref(S)}!${list_col_map[list_name]}$4:${list_col_map[list_name]}${3+LIST_MAX}",
                        allow_blank=True)
    dv.error = "اختر من القائمة أو أضف العنصر في صفحة الإعدادات"
    dv.errorTitle = "قيمة غير موجودة"
    dv.prompt = "اختر من القائمة"
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{DATA_START}:{col_letter}{DATA_END}")


# ==========================================================================
#            محرّك عام لبناء صفحات الإدخال من مواصفات الأعمدة
# ==========================================================================
G_IDENT  = NAVY       # التعريف
G_VEH    = NAVY_MID   # السيارة
G_REV    = GREEN      # التحصيل/الإيرادات
G_EXP    = RED        # المصروفات
G_RESULT = GOLD       # النتيجة

def build_entry_sheet(name, subtitle, columns, formula_fn=None,
                      validations=None, sum_cols=None, key_col_index=2,
                      widths=None):
    """
    columns: list of dict(header, group, kind, numfmt)
    formula_fn(row, L) -> dict{col_index_1based: formula_str}   (L = {header: letter})
    validations: dict{col_index_1based: list_name}
    sum_cols: list of col_index_1based to sum in totals row
    key_col_index: العمود الذي يحدّد إن كان الصف مستخدماً (عادة التاريخ)
    """
    ws = wb.create_sheet(title=name)
    ncols = len(columns)
    headers = [c["header"] for c in columns]
    gcolors = [c["group"] for c in columns]
    banner(ws, ncols, name, subtitle)
    style_headers(ws, headers, gcolors)

    L = {c["header"]: get_column_letter(i) for i, c in enumerate(columns, start=1)}
    key_letter = get_column_letter(key_col_index)

    for i, c in enumerate(columns, start=1):
        letter = get_column_letter(i)
        for r in range(DATA_START, DATA_END + 1):
            cell = ws.cell(row=r, column=i)
            style_data_cell(cell, "input" if c["kind"] == "input" else "formula",
                            c.get("numfmt"), zebra=(r % 2 == 0))
            if c["kind"] == "serial":
                cell.value = (f'=IF(${key_letter}{r}="","",'
                              f'COUNTA(${key_letter}${DATA_START}:${key_letter}{r}))')
            elif c["kind"] == "formula" and formula_fn:
                f = formula_fn(r, L)
                if i in f:
                    cell.value = f'=IF(${key_letter}{r}="","",{f[i]})'

    # القوائم المنسدلة
    if validations:
        for idx, list_name in validations.items():
            add_list_validation(ws, get_column_letter(idx), list_name)

    # الأعمدة
    if widths:
        set_widths(ws, widths)
    else:
        set_widths(ws, [12] * ncols)

    # صف الإجماليات
    if sum_cols:
        totals_row(ws, ncols, [(c, columns[c-1].get("numfmt", MONEY)) for c in sum_cols])

    finalize_entry_sheet(ws, ncols)
    return ws, L


# ==========================================================================
#                        (2) صفحة البوالص
# ==========================================================================
bawalis_cols = [
    dict(header="م", group=G_IDENT, kind="serial", numfmt=INTF),
    dict(header="التاريخ", group=G_IDENT, kind="input", numfmt=DATEF),
    dict(header="رقم البوليصة", group=G_IDENT, kind="input", numfmt=None),
    dict(header="صاحب البيان", group=G_IDENT, kind="input", numfmt=None),
    dict(header="التصنيف", group=G_IDENT, kind="input", numfmt=None),
    dict(header="اسم الشركة", group=G_IDENT, kind="input", numfmt=None),
    dict(header="الجهة", group=G_IDENT, kind="input", numfmt=None),
    dict(header="رقم الحاوية", group=G_IDENT, kind="input", numfmt=None),
    dict(header="رقم السيارة", group=G_VEH, kind="input", numfmt=None),
    dict(header="اسم السائق", group=G_VEH, kind="input", numfmt=None),
    dict(header="صاحب السيارة", group=G_VEH, kind="input", numfmt=None),
    dict(header="نولون", group=G_REV, kind="input", numfmt=MONEY),
    dict(header="عهدة", group=G_REV, kind="input", numfmt=MONEY),
    dict(header="تحويلة", group=G_REV, kind="input", numfmt=MONEY),
    dict(header="خصم", group=G_REV, kind="input", numfmt=MONEY),
    dict(header="بياتة (تحصيل)", group=G_REV, kind="input", numfmt=MONEY),
    dict(header="دفعة", group=G_REV, kind="input", numfmt=MONEY),
    dict(header="باقي", group=G_REV, kind="formula", numfmt=MONEY),
    dict(header="إجمالي التحصيل", group=G_REV, kind="formula", numfmt=MONEY),
    dict(header="دور سواق", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="دور تبّاع", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="بياتة (مصروف)", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="مصروف", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="كرتات", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="تحيا مصر", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="ثبتان", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="ظ كاوتش", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="تصاريح", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="كاوتش", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="م أخرى", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="جاز", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="إجمالي المصروف", group=G_EXP, kind="formula", numfmt=MONEY),
    dict(header="الباقي من العهدة", group=G_RESULT, kind="formula", numfmt=MONEY),
    dict(header="صافي السيارة", group=G_RESULT, kind="formula", numfmt=MONEY),
    dict(header="ربح/خسارة %", group=G_RESULT, kind="formula", numfmt=PCT),
]

# خريطة العنوان -> رقم العمود (تتكيّف تلقائياً مع أي إضافة عمود مثل "خصم")
IDX = {c["header"]: i for i, c in enumerate(bawalis_cols, start=1)}

def bawalis_formula(r, L):
    freight = f"${L['نولون']}{r}"
    custody = f"${L['عهدة']}{r}"
    transfer = f"${L['تحويلة']}{r}"
    overn_r = f"${L['بياتة (تحصيل)']}{r}"
    payment = f"${L['دفعة']}{r}"
    total_r = f"${L['إجمالي التحصيل']}{r}"
    total_e = f"${L['إجمالي المصروف']}{r}"
    exp_first = L["دور سواق"]; exp_last = L["جاز"]
    return {
        IDX["باقي"]: f"N({total_r})-N({payment})",
        IDX["إجمالي التحصيل"]: f"N({freight})+N({transfer})+N({overn_r})",
        IDX["إجمالي المصروف"]: f"SUM(${exp_first}{r}:${exp_last}{r})",
        IDX["الباقي من العهدة"]: f"N({custody})-N({total_e})",
        IDX["صافي السيارة"]: f"N({freight})-N({total_e})",
        IDX["ربح/خسارة %"]: f'IF(N({freight})=0,"",(N({freight})-N({total_e}))/N({freight}))',
    }

_sum_headers = ["نولون", "عهدة", "تحويلة", "خصم", "بياتة (تحصيل)", "دفعة", "باقي",
                "إجمالي التحصيل", "دور سواق", "دور تبّاع", "بياتة (مصروف)", "مصروف",
                "كرتات", "تحيا مصر", "ثبتان", "ظ كاوتش", "تصاريح", "كاوتش", "م أخرى",
                "جاز", "إجمالي المصروف", "الباقي من العهدة", "صافي السيارة"]
bawalis_sum_cols = [IDX[h] for h in _sum_headers]

bawalis_widths = [5, 12, 13, 14, 11, 16, 12, 13, 11, 14, 14,
                  12, 11, 11, 11, 12, 11, 11, 13, 11, 11, 12, 11, 10, 11, 10, 10, 11, 10, 10, 10,
                  14, 15, 13, 12]

ws_baw, Lb = build_entry_sheet(
    "البوالص",
    "سجل بوالص الشحن ومحاسبة رحلات السيارات",
    bawalis_cols, bawalis_formula,
    validations={5: "التصنيفات", 6: "الشركات", 7: "الجهات",
                 10: "السائقين", 11: "أصحاب السيارات"},
    sum_cols=bawalis_sum_cols,
    widths=bawalis_widths)

# تنسيق شرطي: ربح/خسارة السيارة
net_col = Lb["صافي السيارة"]
ws_baw.conditional_formatting.add(
    f"{net_col}{DATA_START}:{net_col}{DATA_END}",
    CellIsRule(operator="lessThan", formula=["0"],
               fill=PatternFill("solid", fgColor=RED_LT),
               font=Font(color=RED, bold=True)))
ws_baw.conditional_formatting.add(
    f"{net_col}{DATA_START}:{net_col}{DATA_END}",
    CellIsRule(operator="greaterThan", formula=["0"],
               font=Font(color=GREEN, bold=True)))
pct_col = Lb["ربح/خسارة %"]
ws_baw.conditional_formatting.add(
    f"{pct_col}{DATA_START}:{pct_col}{DATA_END}",
    CellIsRule(operator="lessThan", formula=["0"],
               fill=PatternFill("solid", fgColor=RED_LT), font=Font(color=RED, bold=True)))


# ==========================================================================
#                        (3) صفحة التخليص
# ==========================================================================
takhlis_cols = [
    dict(header="م", group=G_IDENT, kind="serial", numfmt=INTF),
    dict(header="التاريخ", group=G_IDENT, kind="input", numfmt=DATEF),
    dict(header="اسم الشهادة", group=G_IDENT, kind="input", numfmt=None),
    dict(header="رقم الشهادة", group=G_IDENT, kind="input", numfmt=None),
    dict(header="الأتعاب", group=G_REV, kind="input", numfmt=MONEY),
    dict(header="المصروفات", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="إجمالي الكشف", group=G_RESULT, kind="formula", numfmt=MONEY),
    dict(header="الدفعات", group=G_REV, kind="input", numfmt=MONEY),
    dict(header="المتبقي", group=G_RESULT, kind="formula", numfmt=MONEY),
]
def takhlis_formula(r, L):
    fees = f"${L['الأتعاب']}{r}"; exp = f"${L['المصروفات']}{r}"
    tot = f"${L['إجمالي الكشف']}{r}"; pay = f"${L['الدفعات']}{r}"
    return {
        7: f"N({fees})+N({exp})",     # إجمالي الكشف
        9: f"N({tot})-N({pay})",      # المتبقي
    }
ws_tak, Lt = build_entry_sheet(
    "التخليص", "سجل شهادات وأعمال التخليص الجمركي",
    takhlis_cols, takhlis_formula,
    validations=None,
    sum_cols=[5, 6, 7, 8, 9],
    widths=[6, 13, 26, 16, 14, 14, 15, 14, 15])
rem_col = Lt["المتبقي"]
ws_tak.conditional_formatting.add(
    f"{rem_col}{DATA_START}:{rem_col}{DATA_END}",
    CellIsRule(operator="greaterThan", formula=["0"],
               fill=PatternFill("solid", fgColor=GOLD_LIGHT), font=Font(color=RED, bold=True)))


# ==========================================================================
#                   (4) صفحتا المولدات (ايجي كول / الشجيع)
# ==========================================================================
def gen_cols():
    return [
        dict(header="م", group=G_IDENT, kind="serial", numfmt=INTF),
        dict(header="التاريخ", group=G_IDENT, kind="input", numfmt=DATEF),
        dict(header="اسم الشركة", group=G_IDENT, kind="input", numfmt=None),
        dict(header="العنوان", group=G_IDENT, kind="input", numfmt=None),
        dict(header="رقم الإذن", group=G_IDENT, kind="input", numfmt=None),
        dict(header="عدد المولدات", group=G_VEH, kind="input", numfmt=INTF),
        dict(header="النولون", group=G_REV, kind="input", numfmt=MONEY),
        dict(header="الدفعات", group=G_REV, kind="input", numfmt=MONEY),
        dict(header="المتبقي", group=G_RESULT, kind="formula", numfmt=MONEY),
    ]
def gen_formula(r, L):
    fr = f"${L['النولون']}{r}"; pay = f"${L['الدفعات']}{r}"
    return {9: f"N({fr})-N({pay})"}
gen_widths = [6, 13, 22, 26, 14, 12, 14, 14, 15]

ws_egy, Le = build_entry_sheet(
    "مولدات ايجي كول", "سجل نقل مولدات شركة ايجي كول",
    gen_cols(), gen_formula, validations={3: "الشركات"},
    sum_cols=[6, 7, 8, 9], widths=gen_widths)
ws_shg, Lsh = build_entry_sheet(
    "مولدات الشجيع", "سجل نقل مولدات الشجيع",
    gen_cols(), gen_formula, validations={3: "الشركات"},
    sum_cols=[6, 7, 8, 9], widths=gen_widths)
for wsg in (ws_egy, ws_shg):
    rc = get_column_letter(9)
    wsg.conditional_formatting.add(
        f"{rc}{DATA_START}:{rc}{DATA_END}",
        CellIsRule(operator="greaterThan", formula=["0"],
                   fill=PatternFill("solid", fgColor=GOLD_LIGHT), font=Font(color=RED, bold=True)))


# ==========================================================================
#                        (5) صفحة شامل (الفواتير)
# ==========================================================================
shamil_cols = [
    dict(header="مسلسل", group=G_IDENT, kind="serial", numfmt=INTF),
    dict(header="تاريخ الفاتورة", group=G_IDENT, kind="input", numfmt=DATEF),
    dict(header="رقم الكشف", group=G_IDENT, kind="input", numfmt=None),
    dict(header="رقم الفاتورة", group=G_IDENT, kind="input", numfmt=None),
    dict(header="رقم الإذن", group=G_IDENT, kind="input", numfmt=None),
    dict(header="اسم العميل", group=G_IDENT, kind="input", numfmt=None),
    dict(header="اسم الشركة", group=G_IDENT, kind="input", numfmt=None),
    dict(header="الجهة", group=G_IDENT, kind="input", numfmt=None),
    dict(header="عدد الحاويات", group=G_VEH, kind="input", numfmt=INTF),
    dict(header="النولون", group=G_REV, kind="input", numfmt=MONEY),
    dict(header="بيانات أو عطالات", group=G_REV, kind="input", numfmt=MONEY),
    dict(header="قيمة مضافة", group=G_REV, kind="formula", numfmt=MONEY),
    dict(header="خصم أ ت ص", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="الإيصالات", group=G_REV, kind="input", numfmt=MONEY),
    dict(header="إجمالي الفاتورة", group=G_RESULT, kind="formula", numfmt=MONEY),
    dict(header="الدفعات", group=G_REV, kind="input", numfmt=MONEY),
    dict(header="إجمالي الحساب", group=G_RESULT, kind="formula", numfmt=MONEY),
    dict(header="نولون فعلي", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="كشف وضمان", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="مصروف", group=G_EXP, kind="input", numfmt=MONEY),
    dict(header="إجمالي ربح", group=G_RESULT, kind="formula", numfmt=MONEY),
    dict(header="ربح آخر", group=G_RESULT, kind="input", numfmt=MONEY),
]
def shamil_formula(r, L):
    freight = f"${L['النولون']}{r}"
    dataidle = f"${L['بيانات أو عطالات']}{r}"
    vat = f"${L['قيمة مضافة']}{r}"
    disc = f"${L['خصم أ ت ص']}{r}"
    receipts = f"${L['الإيصالات']}{r}"
    inv_total = f"${L['إجمالي الفاتورة']}{r}"
    payments = f"${L['الدفعات']}{r}"
    actual = f"${L['نولون فعلي']}{r}"
    expense = f"${L['مصروف']}{r}"
    return {
        12: f"ROUND((N({freight})+N({dataidle}))*نسبة_القيمة_المضافة,2)",              # قيمة مضافة
        15: f"N({freight})+N({dataidle})+N({vat})-N({disc})+N({receipts})",            # إجمالي الفاتورة
        17: f"N({inv_total})-N({payments})",                                           # إجمالي الحساب
        21: f"N({freight})-N({actual})-N({expense})",                                  # إجمالي ربح
    }
shamil_widths = [7, 13, 11, 11, 11, 18, 16, 12, 11, 12, 13, 12, 11, 11, 14, 12, 14, 12, 12, 11, 12, 11]

ws_sha, Ls = build_entry_sheet(
    "شامل", "السجل الشامل للفواتير وحسابات العملاء",
    shamil_cols, shamil_formula,
    validations={6: "العملاء", 7: "الشركات", 8: "الجهات"},
    sum_cols=[9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22],
    widths=shamil_widths)
acc_col = Ls["إجمالي الحساب"]
ws_sha.conditional_formatting.add(
    f"{acc_col}{DATA_START}:{acc_col}{DATA_END}",
    CellIsRule(operator="greaterThan", formula=["0"],
               fill=PatternFill("solid", fgColor=GOLD_LIGHT), font=Font(color=RED, bold=True)))
prof_col = Ls["إجمالي ربح"]
ws_sha.conditional_formatting.add(
    f"{prof_col}{DATA_START}:{prof_col}{DATA_END}",
    CellIsRule(operator="lessThan", formula=["0"],
               fill=PatternFill("solid", fgColor=RED_LT), font=Font(color=RED, bold=True)))


# ==========================================================================
#              (6) صفحة بيانات التقارير (تجميعات SUMIFS)
# ==========================================================================
def rng(sheet, letter):
    return f"{sref(sheet)}!${letter}${DATA_START}:${letter}${DATA_END}"

# اختصارات الأعمدة
BAW = "البوالص"; TAK = "التخليص"; EGY = "مولدات ايجي كول"; SHG = "مولدات الشجيع"; SHA = "شامل"
b_date = Lb["التاريخ"]; b_freight = Lb["نولون"]; b_company = Lb["اسم الشركة"]
b_net = Lb["صافي السيارة"]; b_texp = Lb["إجمالي المصروف"]; b_container = Lb["رقم الحاوية"]
b_payment = Lb["دفعة"]; b_baqi = Lb["باقي"]
t_date = Lt["التاريخ"]; t_tot = Lt["إجمالي الكشف"]; t_pay = Lt["الدفعات"]; t_rem = Lt["المتبقي"]; t_exp = Lt["المصروفات"]; t_fees = Lt["الأتعاب"]
e_date = Le["التاريخ"]; e_freight = Le["النولون"]; e_pay = Le["الدفعات"]; e_rem = Le["المتبقي"]; e_cnt = Le["عدد المولدات"]; e_comp = Le["اسم الشركة"]
sh_date = Ls["تاريخ الفاتورة"]; sh_client = Ls["اسم العميل"]; sh_comp = Ls["اسم الشركة"]
sh_inv = Ls["إجمالي الفاتورة"]; sh_pay = Ls["الدفعات"]; sh_acc = Ls["إجمالي الحساب"]
sh_profit = Ls["إجمالي ربح"]; sh_other = Ls["ربح آخر"]; sh_freight = Ls["النولون"]
sh_cont = Ls["عدد الحاويات"]; sh_exp = Ls["مصروف"]

RP = "بيانات_التقارير"
ws_rp = wb.create_sheet(title=RP)
ws_rp.sheet_view.rightToLeft = True
ws_rp.sheet_view.showGridLines = False
banner(ws_rp, 6, "بيانات التقارير (تُحدَّث تلقائياً)", "لا تُعدّل هذه الصفحة يدوياً — تُغذّي لوحة التحكم والرسوم")

def hdr(ws, cell, text, color=NAVY):
    c = ws[cell]; c.value = text
    c.font = Font(name=FONT_NAME, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=color)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = box_thin

def money_cell(ws, coord, formula):
    c = ws[coord]; c.value = formula
    c.number_format = MONEY0
    c.font = Font(name=FONT_NAME, color=INK)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = box_thin

def month_sumif(vs, vl, ds, dl, mrow):
    return (f'SUMIFS({rng(vs,vl)},{rng(ds,dl)},">="&$A{mrow},'
            f'{rng(ds,dl)},"<"&EDATE($A{mrow},1))')

# ---- جدول شهري (الصفوف 4..15) ----
MONTHS = ["يناير","فبراير","مارس","أبريل","مايو","يونيو","يوليو","أغسطس","سبتمبر","أكتوبر","نوفمبر","ديسمبر"]
hdr(ws_rp, "A3", "تاريخ"); hdr(ws_rp, "B3", "الشهر"); hdr(ws_rp, "C3", "الإيرادات", GREEN)
hdr(ws_rp, "D3", "المصروفات", RED); hdr(ws_rp, "E3", "صافي الربح", GOLD); hdr(ws_rp, "F3", "المتحصّلات", TEAL)
for m in range(1, 13):
    r = 3 + m
    ws_rp[f"A{r}"] = f"=DATE(سنة_التقارير,{m},1)"
    ws_rp[f"A{r}"].number_format = "MMM-YYYY"
    ws_rp[f"A{r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws_rp[f"A{r}"].border = box_thin
    bc = ws_rp[f"B{r}"]; bc.value = MONTHS[m-1]
    bc.font = Font(name=FONT_NAME, bold=True, color=INK); bc.border = box_thin
    bc.alignment = Alignment(horizontal="center", vertical="center")
    bc.fill = PatternFill("solid", fgColor=SECTION_BG)
    rev = "+".join([month_sumif(BAW, b_freight, BAW, b_date, r),
                    month_sumif(SHA, sh_inv, SHA, sh_date, r),
                    month_sumif(TAK, t_tot, TAK, t_date, r),
                    month_sumif(EGY, e_freight, EGY, e_date, r),
                    month_sumif(SHG, e_freight, SHG, e_date, r)])
    exp = "+".join([month_sumif(BAW, b_texp, BAW, b_date, r),
                    month_sumif(SHA, sh_exp, SHA, sh_date, r),
                    month_sumif(TAK, t_exp, TAK, t_date, r)])
    prof = "+".join([month_sumif(SHA, sh_profit, SHA, sh_date, r),
                     month_sumif(SHA, sh_other, SHA, sh_date, r),
                     month_sumif(BAW, b_net, BAW, b_date, r)])
    coll = "+".join([month_sumif(BAW, b_payment, BAW, b_date, r),
                     month_sumif(SHA, sh_pay, SHA, sh_date, r),
                     month_sumif(TAK, t_pay, TAK, t_date, r),
                     month_sumif(EGY, e_pay, EGY, e_date, r),
                     month_sumif(SHG, e_pay, SHG, e_date, r)])
    money_cell(ws_rp, f"C{r}", "=" + rev)
    money_cell(ws_rp, f"D{r}", "=" + exp)
    money_cell(ws_rp, f"E{r}", "=" + prof)
    money_cell(ws_rp, f"F{r}", "=" + coll)

# ---- حسب الشركة (الأعمدة H..J، الصفوف 4..18) ----
NLIST = 15
hdr(ws_rp, "H3", "الشركة"); hdr(ws_rp, "I3", "النولون", GREEN); hdr(ws_rp, "J3", "صافي الربح", GOLD)
for i in range(NLIST):
    r = 4 + i
    ws_rp[f"H{r}"] = f"=IF(الإعدادات!$E{4+i}=\"\",\"\",الإعدادات!$E{4+i})"
    ws_rp[f"H{r}"].font = Font(name=FONT_NAME, color=INK); ws_rp[f"H{r}"].border = box_thin
    ws_rp[f"H{r}"].alignment = Alignment(horizontal="center", vertical="center")
    cc = f"الإعدادات!$E{4+i}"
    money_cell(ws_rp, f"I{r}",
               f'=SUMIF({rng(BAW,b_company)},{cc},{rng(BAW,b_freight)})+SUMIF({rng(SHA,sh_comp)},{cc},{rng(SHA,sh_freight)})')
    money_cell(ws_rp, f"J{r}",
               f'=SUMIF({rng(BAW,b_company)},{cc},{rng(BAW,b_net)})+SUMIF({rng(SHA,sh_comp)},{cc},{rng(SHA,sh_profit)})')

# ---- أعلى العملاء (الأعمدة L..N) ----
hdr(ws_rp, "L3", "العميل"); hdr(ws_rp, "M3", "إجمالي الفواتير", GREEN); hdr(ws_rp, "N3", "المتبقي عليه", RED)
for i in range(NLIST):
    r = 4 + i
    ws_rp[f"L{r}"] = f"=IF(الإعدادات!$I{4+i}=\"\",\"\",الإعدادات!$I{4+i})"
    ws_rp[f"L{r}"].font = Font(name=FONT_NAME, color=INK); ws_rp[f"L{r}"].border = box_thin
    ws_rp[f"L{r}"].alignment = Alignment(horizontal="center", vertical="center")
    cc = f"الإعدادات!$I{4+i}"
    money_cell(ws_rp, f"M{r}", f'=SUMIF({rng(SHA,sh_client)},{cc},{rng(SHA,sh_inv)})')
    money_cell(ws_rp, f"N{r}", f'=SUMIF({rng(SHA,sh_client)},{cc},{rng(SHA,sh_acc)})')

# ---- توزيع مصروفات البوالص (الأعمدة P..Q) ----
exp_names = ["دور سواق","دور تبّاع","بياتة (مصروف)","مصروف","كرتات","تحيا مصر","ثبتان","ظ كاوتش","تصاريح","كاوتش","م أخرى","جاز"]
hdr(ws_rp, "P3", "بند المصروف"); hdr(ws_rp, "Q3", "الإجمالي", RED)
for i, en in enumerate(exp_names):
    r = 4 + i
    pc = ws_rp[f"P{r}"]; pc.value = en
    pc.font = Font(name=FONT_NAME, color=INK); pc.border = box_thin
    pc.alignment = Alignment(horizontal="center", vertical="center")
    pc.fill = PatternFill("solid", fgColor=SECTION_BG)
    money_cell(ws_rp, f"Q{r}", f"=SUM({rng(BAW, Lb[en])})")

# ---- مقارنة المولدات (الأعمدة S..T) ----
hdr(ws_rp, "S3", "جهة المولدات"); hdr(ws_rp, "T3", "إجمالي النولون", TEAL)
ws_rp["S4"] = "ايجي كول"; ws_rp["S5"] = "الشجيع"
for r in (4, 5):
    ws_rp[f"S{r}"].font = Font(name=FONT_NAME, bold=True, color=INK)
    ws_rp[f"S{r}"].border = box_thin
    ws_rp[f"S{r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws_rp[f"S{r}"].fill = PatternFill("solid", fgColor=SECTION_BG)
money_cell(ws_rp, "T4", f"=SUM({rng(EGY, e_freight)})")
money_cell(ws_rp, "T5", f"=SUM({rng(SHG, e_freight)})")

# ---- مرآة أرصدة الموردين (تُقرأ من صفحة الأرصدة) للرسم البياني ----
AR_NAME = "أرصدة العملاء والموردين"
hdr(ws_rp, "V3", "المورد"); hdr(ws_rp, "W3", "الرصيد", RED)
for i in range(12):
    rr = 4 + i
    vc = ws_rp[f"V{rr}"]
    vc.value = f"=IF('{AR_NAME}'!E{4+i}=\"\",\"\",'{AR_NAME}'!E{4+i})"
    vc.font = Font(name=FONT_NAME, color=INK); vc.border = box_thin
    vc.alignment = Alignment(horizontal="center", vertical="center")
    money_cell(ws_rp, f"W{rr}", f"='{AR_NAME}'!F{4+i}")

set_widths(ws_rp, [12, 12, 15, 15, 15, 15, 2, 18, 15, 15, 2, 18, 16, 15, 2, 16, 15, 2, 14, 16, 2, 18, 14])


# ==========================================================================
#                        (7) لوحة التحكم (Dashboard)
# ==========================================================================
def sumall(sheet, letter):
    return f"SUM({rng(sheet, letter)})"

DB = "لوحة التحكم"
ws_db = wb.create_sheet(title=DB)
ws_db.sheet_view.rightToLeft = True
ws_db.sheet_view.showGridLines = False
NC = 12
banner(ws_db, NC, "لوحة التحكم والمؤشرات", "نظرة شاملة على أداء الشركة — تُحدَّث تلقائياً مع إدخال البيانات")
set_widths(ws_db, [11] * NC)
ws_db.row_dimensions[3].height = 26   # شريط أزرار الماكرو (يملؤه الماكرو)

def fill_block(ws, r1, c1, r2, c2, color):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=color)

def border_block(ws, r1, c1, r2, c2, side):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = Border(left=side, right=side, top=side, bottom=side)

def make_card(ws, r1, c1, r2, c2, label, formula, numfmt, accent):
    # شريط العنوان
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r1, end_column=c2)
    fill_block(ws, r1, c1, r1, c1_ := c2, accent)
    lc = ws.cell(r1, c1, label)
    lc.font = Font(name=FONT_NAME, bold=True, color=WHITE, size=11)
    lc.alignment = Alignment(horizontal="center", vertical="center")
    # القيمة
    ws.merge_cells(start_row=r1 + 1, start_column=c1, end_row=r2, end_column=c2)
    fill_block(ws, r1 + 1, c1, r2, c2, WHITE)
    vc = ws.cell(r1 + 1, c1, formula)
    vc.number_format = numfmt
    vc.font = Font(name=FONT_NAME, bold=True, size=18, color=accent)
    vc.alignment = Alignment(horizontal="center", vertical="center")
    border_block(ws, r1, c1, r2, c2, Side(style="thin", color=accent))
    ws.row_dimensions[r1].height = 20
    for rr in range(r1 + 1, r2 + 1):
        ws.row_dimensions[rr].height = 22

rev_f = "=" + "+".join([sumall(BAW, b_freight), sumall(SHA, sh_inv), sumall(TAK, t_tot), sumall(EGY, e_freight), sumall(SHG, e_freight)])
exp_f = "=" + "+".join([sumall(BAW, b_texp), sumall(SHA, sh_exp), sumall(TAK, t_exp)])
prof_f = "=" + "+".join([sumall(SHA, sh_profit), sumall(SHA, sh_other), sumall(BAW, b_net)])
coll_f = "=" + "+".join([sumall(BAW, b_payment), sumall(SHA, sh_pay), sumall(TAK, t_pay), sumall(EGY, e_pay), sumall(SHG, e_pay)])
out_f = "=" + "+".join([sumall(BAW, b_baqi), sumall(SHA, sh_acc), sumall(TAK, t_rem), sumall(EGY, e_rem), sumall(SHG, e_rem)])
freight_f = "=" + "+".join([sumall(BAW, b_freight), sumall(SHA, sh_freight), sumall(EGY, e_freight), sumall(SHG, e_freight)])

# الصف الأول من البطاقات (rows 4-6)
make_card(ws_db, 4, 1, 6, 4, "إجمالي الإيرادات", rev_f, MONEY0, GREEN)
make_card(ws_db, 4, 5, 6, 8, "إجمالي المصروفات", exp_f, MONEY0, RED)
make_card(ws_db, 4, 9, 6, 12, "صافي الأرباح", prof_f, MONEY0, GOLD)
# الصف الثاني (rows 7-9)
make_card(ws_db, 7, 1, 9, 4, "إجمالي المتحصّلات", coll_f, MONEY0, TEAL)
make_card(ws_db, 7, 5, 9, 8, "المتبقي على العملاء", out_f, MONEY0, PURPLE)
make_card(ws_db, 7, 9, 9, 12, "إجمالي النولون", freight_f, MONEY0, NAVY_MID)

# صف عدّادات صغيرة (rows 10-11)
cnt_baw = f"=COUNT({rng(BAW, b_date)})"
cnt_inv = f"=COUNT({rng(SHA, sh_date)})"
cnt_cont = f"=SUM({rng(SHA, sh_cont)})+SUMPRODUCT(--({rng(BAW, b_container)}<>\"\"))"
cnt_gen = f"=SUM({rng(EGY, e_cnt)})+SUM({rng(SHG, e_cnt)})"
cnt_tak = f"=COUNT({rng(TAK, t_date)})"
small = [("عدد البوالص", cnt_baw), ("عدد الفواتير", cnt_inv), ("عدد الحاويات", cnt_cont),
         ("عدد المولدات", cnt_gen), ("شهادات التخليص", cnt_tak)]
# 5 عدّادات بعرضين
sc = 1
for i, (lab, f) in enumerate(small):
    c1 = 1 + i * 2 + (1 if i >= 5 else 0)
    c2 = c1 + 1
    if c2 > 12:
        break
    ws_db.merge_cells(start_row=10, start_column=c1, end_row=10, end_column=c2)
    ws_db.merge_cells(start_row=11, start_column=c1, end_row=11, end_column=c2)
    fill_block(ws_db, 10, c1, 10, c2, SECTION_BG)
    fill_block(ws_db, 11, c1, 11, c2, WHITE)
    lc = ws_db.cell(10, c1, lab)
    lc.font = Font(name=FONT_NAME, bold=True, color=NAVY, size=9)
    lc.alignment = Alignment(horizontal="center", vertical="center")
    vc = ws_db.cell(11, c1, f)
    vc.number_format = INTF
    vc.font = Font(name=FONT_NAME, bold=True, color=INK, size=14)
    vc.alignment = Alignment(horizontal="center", vertical="center")
    border_block(ws_db, 10, c1, 11, c2, Side(style="thin", color=GRID))
ws_db.row_dimensions[10].height = 16
ws_db.row_dimensions[11].height = 22

# عنوان قسم الرسوم
ws_db.merge_cells("A13:L13")
tt = ws_db["A13"]; tt.value = "التحليلات البيانية"
tt.font = Font(name=FONT_NAME, bold=True, size=13, color=WHITE)
tt.alignment = Alignment(horizontal="center", vertical="center")
fill_block(ws_db, 13, 1, 13, 12, NAVY)
ws_db.row_dimensions[13].height = 24

# ---- الرسوم البيانية ----
CHW, CHH = 15, 8

# 1) الأداء الشهري (خطي)
lc1 = LineChart(); lc1.title = "الأداء الشهري (إيرادات / مصروفات / أرباح)"
data = Reference(ws_rp, min_col=3, max_col=5, min_row=3, max_row=15)
cats = Reference(ws_rp, min_col=2, max_col=2, min_row=4, max_row=15)
lc1.add_data(data, titles_from_data=True); lc1.set_categories(cats)
lc1.width, lc1.height = CHW, CHH
lc1.style = 12
ws_db.add_chart(lc1, "A15")

# 2) صافي الربح حسب الشركة (أعمدة)
bc1 = BarChart(); bc1.type = "bar"; bc1.title = "صافي الربح حسب الشركة"
data = Reference(ws_rp, min_col=10, max_col=10, min_row=3, max_row=18)
cats = Reference(ws_rp, min_col=8, max_col=8, min_row=4, max_row=18)
bc1.add_data(data, titles_from_data=True); bc1.set_categories(cats)
bc1.width, bc1.height = CHW, CHH; bc1.style = 10
ws_db.add_chart(bc1, "G15")

# 3) توزيع المصروفات (دائري)
pc1 = PieChart(); pc1.title = "توزيع مصروفات البوالص"
data = Reference(ws_rp, min_col=17, max_col=17, min_row=3, max_row=15)
cats = Reference(ws_rp, min_col=16, max_col=16, min_row=4, max_row=15)
pc1.add_data(data, titles_from_data=True); pc1.set_categories(cats)
pc1.width, pc1.height = CHW, CHH
pc1.dataLabels = DataLabelList(); pc1.dataLabels.showPercent = True
ws_db.add_chart(pc1, "A32")

# 4) أرصدة الموردين (أعمدة)
bc2 = BarChart(); bc2.type = "bar"; bc2.title = "أرصدة الموردين"
data = Reference(ws_rp, min_col=23, max_col=23, min_row=3, max_row=15)
cats = Reference(ws_rp, min_col=22, max_col=22, min_row=4, max_row=15)
bc2.add_data(data, titles_from_data=True); bc2.set_categories(cats)
bc2.width, bc2.height = CHW, CHH; bc2.style = 12
ws_db.add_chart(bc2, "G32")

# 5) مقارنة المولدات (دائري)
pc2 = PieChart(); pc2.title = "مقارنة نولون المولدات"
data = Reference(ws_rp, min_col=20, max_col=20, min_row=3, max_row=5)
cats = Reference(ws_rp, min_col=19, max_col=19, min_row=4, max_row=5)
pc2.add_data(data, titles_from_data=True); pc2.set_categories(cats)
pc2.width, pc2.height = CHW, CHH
pc2.dataLabels = DataLabelList(); pc2.dataLabels.showVal = True
ws_db.add_chart(pc2, "A49")

# ضمان ظهور بيانات الرسوم رغم أن صفحة المصدر مخفية
for _ch in (lc1, bc1, pc1, bc2, pc2):
    _ch.visible_cells_only = False
    _ch.display_blanks = "gap"

ws_db.page_setup.orientation = "landscape"
ws_db.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


# ==========================================================================
#                        (8) صفحة التقارير
# ==========================================================================
RPT = "التقارير"
ws_r = wb.create_sheet(title=RPT)
ws_r.sheet_view.rightToLeft = True
ws_r.sheet_view.showGridLines = False
banner(ws_r, 10, "التقارير المالية والإدارية", "تقارير جاهزة للطباعة — تُحدَّث تلقائياً")
set_widths(ws_r, [4, 20, 18, 4, 20, 18, 4, 20, 18, 4])

def section_title(ws, row, first_col, last_col, text, color=NAVY):
    ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)
    fill_block(ws, row, first_col, row, last_col, color)
    c = ws.cell(row, first_col, text)
    c.font = Font(name=FONT_NAME, bold=True, size=12, color=WHITE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 24

def kv_row(ws, row, col_label, col_val, label, formula, accent, numfmt=MONEY0):
    lc = ws.cell(row, col_label, label)
    lc.font = Font(name=FONT_NAME, bold=True, color=INK)
    lc.fill = PatternFill("solid", fgColor=SECTION_BG)
    lc.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    lc.border = box_thin
    vc = ws.cell(row, col_val, formula)
    vc.number_format = numfmt
    vc.font = Font(name=FONT_NAME, bold=True, color=accent, size=12)
    vc.alignment = Alignment(horizontal="center", vertical="center")
    vc.fill = PatternFill("solid", fgColor=WHITE)
    vc.border = box_thin
    ws.row_dimensions[row].height = 20

# --- الملخص المالي العام ---
section_title(ws_r, 3, 2, 9, "الملخص المالي العام")
summary = [
    ("إجمالي الإيرادات", rev_f, GREEN),
    ("إجمالي المصروفات", exp_f, RED),
    ("صافي الأرباح", prof_f, GOLD),
    ("إجمالي المتحصّلات", coll_f, TEAL),
    ("المتبقي على العملاء", out_f, PURPLE),
    ("إجمالي النولون", freight_f, NAVY_MID),
]
for i, (lab, f, acc) in enumerate(summary):
    row = 4 + (i % 3)
    if i < 3:
        kv_row(ws_r, row, 2, 3, lab, f, acc)
    else:
        kv_row(ws_r, row, 5, 6, lab, f, acc)
# نسبة هامش الربح
kv_row(ws_r, 7, 5, 6, "هامش الربح %",
       f'=IF(({rev_f[1:]})=0,"",({prof_f[1:]})/({rev_f[1:]}))', GOLD, PCT)

# --- التقرير الشهري ---
section_title(ws_r, 9, 2, 6, "الأداء الشهري", GREEN)
mh = ["الشهر", "الإيرادات", "المصروفات", "صافي الربح", "المتحصّلات"]
for j, h in enumerate(mh):
    c = ws_r.cell(10, 2 + j, h)
    c.font = Font(name=FONT_NAME, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY_MID)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = box_thin
# ملاحظة: الجدول الشهري 5 أعمدة يحتاج B..F
set_widths(ws_r, [4, 14, 15, 15, 15, 15, 4, 20, 16, 16])
src_cols = ["B", "C", "D", "E", "F"]   # في بيانات_التقارير
for m in range(12):
    row = 11 + m
    for j, sc in enumerate(src_cols):
        c = ws_r.cell(row, 2 + j, f"={sref(RP)}!{sc}{4+m}")
        c.border = box_thin
        c.alignment = Alignment(horizontal="center", vertical="center")
        if j == 0:
            c.font = Font(name=FONT_NAME, bold=True, color=INK)
            c.fill = PatternFill("solid", fgColor=(ZEBRA if m % 2 else WHITE))
        else:
            c.number_format = MONEY0
            c.font = Font(name=FONT_NAME, color=INK)
            c.fill = PatternFill("solid", fgColor=(ZEBRA if m % 2 else WHITE))
# صف إجمالي شهري
trow = 23
for j in range(5):
    c = ws_r.cell(trow, 2 + j)
    c.fill = PatternFill("solid", fgColor=NAVY); c.border = box_thin
    c.font = Font(name=FONT_NAME, bold=True, color=GOLD)
    c.alignment = Alignment(horizontal="center", vertical="center")
ws_r.cell(trow, 2, "الإجمالي السنوي")
for j, sc in enumerate(["C", "D", "E", "F"], start=1):
    c = ws_r.cell(trow, 2 + j, f"=SUM({sref(RP)}!{sc}4:{sc}15)")
    c.number_format = MONEY0

# --- تقرير حسب الشركة (يمين) ---
section_title(ws_r, 9, 8, 10, "الربح حسب الشركة", TEAL)
for j, h in enumerate(["الشركة", "النولون", "صافي الربح"]):
    c = ws_r.cell(10, 8 + j, h)
    c.font = Font(name=FONT_NAME, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY_MID)
    c.alignment = Alignment(horizontal="center", vertical="center"); c.border = box_thin
for i in range(15):
    row = 11 + i
    for j, sc in enumerate(["H", "I", "J"]):
        c = ws_r.cell(row, 8 + j, f"={sref(RP)}!{sc}{4+i}")
        c.border = box_thin; c.alignment = Alignment(horizontal="center", vertical="center")
        c.fill = PatternFill("solid", fgColor=(ZEBRA if i % 2 else WHITE))
        if j == 0:
            c.font = Font(name=FONT_NAME, bold=True, color=INK)
        else:
            c.number_format = MONEY0; c.font = Font(name=FONT_NAME, color=INK)

# --- تقرير مديونية العملاء ---
section_title(ws_r, 26, 2, 6, "مديونية العملاء (المتأخرات)", RED)
for j, h in enumerate(["العميل", "إجمالي الفواتير", "المتبقي عليه"]):
    c = ws_r.cell(27, 2 + j, h)
    c.font = Font(name=FONT_NAME, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY_MID)
    c.alignment = Alignment(horizontal="center", vertical="center"); c.border = box_thin
for i in range(15):
    row = 28 + i
    for j, sc in enumerate(["L", "M", "N"]):
        c = ws_r.cell(row, 2 + j, f"={sref(RP)}!{sc}{4+i}")
        c.border = box_thin; c.alignment = Alignment(horizontal="center", vertical="center")
        c.fill = PatternFill("solid", fgColor=(ZEBRA if i % 2 else WHITE))
        if j == 0:
            c.font = Font(name=FONT_NAME, bold=True, color=INK)
        else:
            c.number_format = MONEY0; c.font = Font(name=FONT_NAME, color=INK)
# إبراز المتبقي > 0
ws_r.conditional_formatting.add(
    f"D28:D42",
    CellIsRule(operator="greaterThan", formula=["0"],
               fill=PatternFill("solid", fgColor=RED_LT), font=Font(color=RED, bold=True)))

ws_r.freeze_panes = "A3"
ws_r.page_setup.orientation = "landscape"
ws_r.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


# ==========================================================================
#                    (9) صفحة الدليل والتعليمات
# ==========================================================================
GD = "الدليل والتعليمات"
ws_g = wb.create_sheet(title=GD)
ws_g.sheet_view.rightToLeft = True
ws_g.sheet_view.showGridLines = False
banner(ws_g, 8, "دليل استخدام النظام", "اقرأ هذه الصفحة قبل البدء")
set_widths(ws_g, [3, 26, 40, 22, 3, 3, 3, 3])

def g_title(row, text, color=NAVY):
    ws_g.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    fill_block(ws_g, row, 2, row, 4, color)
    c = ws_g.cell(row, 2, text)
    c.font = Font(name=FONT_NAME, bold=True, size=12, color=WHITE)
    c.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws_g.row_dimensions[row].height = 22

def g_line(row, a, b, c="", bold=False):
    x = ws_g.cell(row, 2, a)
    x.font = Font(name=FONT_NAME, bold=True, color=NAVY, size=10)
    x.alignment = Alignment(horizontal="right", vertical="center", indent=1, wrap_text=True)
    x.border = box_thin; x.fill = PatternFill("solid", fgColor=SECTION_BG)
    y = ws_g.cell(row, 3, b)
    y.font = Font(name=FONT_NAME, color=INK, size=10, bold=bold)
    y.alignment = Alignment(horizontal="right", vertical="center", indent=1, wrap_text=True)
    y.border = box_thin
    z = ws_g.cell(row, 4, c)
    z.font = Font(name=FONT_NAME, color=GREY_TXT, size=9)
    z.alignment = Alignment(horizontal="right", vertical="center", indent=1, wrap_text=True)
    z.border = box_thin
    ws_g.row_dimensions[row].height = 30

r = 3
g_title(r, "كيف يعمل النظام؟"); r += 1
for a, b, c in [
    ("الفكرة", "تُدخِل بياناتك في صفحات (البوالص، التخليص، المولدات، شامل) فقط، وكل الحسابات والتقارير واللوحة تتحدّث تلقائياً.", ""),
    ("الخلايا الصفراء", "خلايا إدخال — اكتب فيها بياناتك.", "لونها أصفر فاتح"),
    ("الخلايا الزرقاء", "خلايا محسوبة تلقائياً بمعادلات — لا تكتب فيها (النص مائل).", "لونها أزرق فاتح"),
    ("القوائم المنسدلة", "التصنيف/الشركة/الجهة/السائق/العميل تُختار من قائمة. لإضافة عنصر جديد افتح صفحة (الإعدادات).", ""),
    ("الترقيم", "عمود (م) يرقّم تلقائياً بمجرد كتابة التاريخ في الصف.", ""),
]:
    g_line(r, a, b, c); r += 1

r += 1
g_title(r, "شرح المعادلات (يمكنك تعديلها لو نظامك يحسب بطريقة مختلفة)"); r += 1
for a, b, c in [
    ("البوالص › إجمالي التحصيل", "= نولون + تحويلة + بياتة (تحصيل)", "راجعها حسب نظامك"),
    ("البوالص › باقي", "= إجمالي التحصيل − دفعة", ""),
    ("البوالص › إجمالي المصروف", "= مجموع بنود المصروفات (دور سواق ← جاز)", "مؤكّدة"),
    ("البوالص › الباقي من العهدة", "= عهدة − إجمالي المصروف", "مؤكّدة"),
    ("البوالص › صافي السيارة", "= نولون − إجمالي المصروف", "مؤكّدة"),
    ("البوالص › ربح/خسارة %", "= صافي السيارة ÷ نولون", ""),
    ("التخليص › إجمالي الكشف", "= الأتعاب + المصروفات", "مؤكّدة"),
    ("التخليص › المتبقي", "= إجمالي الكشف − الدفعات", "مؤكّدة"),
    ("المولدات › المتبقي", "= النولون − الدفعات", "مؤكّدة"),
    ("شامل › قيمة مضافة", "= (النولون + بيانات/عطالات) × نسبة القيمة المضافة", "النسبة في الإعدادات (14%)"),
    ("شامل › إجمالي الفاتورة", "= نولون + بيانات + ق.مضافة − خصم أ ت ص + الإيصالات", "راجعها حسب نظامك"),
    ("شامل › إجمالي الحساب", "= إجمالي الفاتورة − الدفعات", "المتبقي على العميل"),
    ("شامل › إجمالي ربح", "= النولون − نولون فعلي − مصروف", "راجعها حسب نظامك"),
]:
    g_line(r, a, b, c); r += 1

r += 1
g_title(r, "تفعيل الماكرو (اختياري)", GOLD); r += 1
for a, b, c in [
    ("لماذا؟", "الماكرو يضيف أزرار: فورم إدخال سريع، بحث، طباعة PDF، نسخة احتياطية.", ""),
    ("الخطوة 1", "احفظ الملف بصيغة Excel Macro-Enabled (.xlsm).", "ملف ← حفظ باسم"),
    ("الخطوة 2", "افتح محرر الأكواد بضغط Alt+F11، ثم File ← Import File واختر الملف (وحدة_الماكرو.bas).", ""),
    ("الخطوة 3", "ارجع للإكسيل وشغّل الماكرو من View ← Macros أو من الأزرار.", ""),
    ("ملاحظة", "لو مش عايز ماكرو، النظام يعمل بالكامل بدونه بالمعادلات والرسوم.", ""),
]:
    g_line(r, a, b, c); r += 1

ws_g.freeze_panes = "A3"


# ==========================================================================
#            (9.5) صفحة أرصدة العملاء والموردين (كشف حساب)
# ==========================================================================
AR = "أرصدة العملاء والموردين"
ws_ar = wb.create_sheet(title=AR)
ws_ar.sheet_view.rightToLeft = True
ws_ar.sheet_view.showGridLines = False
banner(ws_ar, 7, "أرصدة العملاء والموردين", "كشف حساب العملاء والموردين — المنقول من النظام القديم")
set_widths(ws_ar, [4, 28, 18, 3, 28, 18, 4])
AR_ROWS = 200
AR_START = 4
AR_END = AR_START + AR_ROWS - 1

def ar_hdr(cell, text, color):
    c = ws_ar[cell]; c.value = text
    c.font = Font(name=FONT_NAME, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=color)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = box_thin
ar_hdr("B3", "العميل", NAVY); ar_hdr("C3", "الرصيد (له)", GREEN)
ar_hdr("E3", "المورد", NAVY); ar_hdr("F3", "الرصيد (عليه)", RED)
ws_ar.row_dimensions[3].height = 26

for r in range(AR_START, AR_END + 1):
    for col, is_money in (("B", False), ("C", True), ("E", False), ("F", True)):
        cell = ws_ar[f"{col}{r}"]
        cell.fill = PatternFill("solid", fgColor=(INPUT_FILL if r % 2 else ZEBRA))
        cell.border = box_thin
        cell.font = Font(name=FONT_NAME, size=10, color=INK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if is_money:
            cell.number_format = MONEY0

# صف الإجماليات
tr = AR_END + 1
for col in ("B", "C", "E", "F"):
    c = ws_ar[f"{col}{tr}"]
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.font = Font(name=FONT_NAME, bold=True, color=GOLD)
    c.border = Border(top=Side(style="double", color=GOLD))
    c.alignment = Alignment(horizontal="center", vertical="center")
ws_ar[f"B{tr}"] = "إجمالي أرصدة العملاء"
ws_ar[f"C{tr}"] = f"=SUM(C{AR_START}:C{AR_END})"; ws_ar[f"C{tr}"].number_format = MONEY0
ws_ar[f"E{tr}"] = "إجمالي أرصدة الموردين"
ws_ar[f"F{tr}"] = f"=SUM(F{AR_START}:F{AR_END})"; ws_ar[f"F{tr}"].number_format = MONEY0
dr = tr + 1
for col in ("E", "F"):
    c = ws_ar[f"{col}{dr}"]
    c.fill = PatternFill("solid", fgColor=GOLD_LIGHT)
    c.font = Font(name=FONT_NAME, bold=True, color=NAVY)
    c.border = box_thin
    c.alignment = Alignment(horizontal="center", vertical="center")
ws_ar[f"E{dr}"] = "الفرق (له − عليه)"
ws_ar[f"F{dr}"] = f"=C{tr}-F{tr}"; ws_ar[f"F{dr}"].number_format = MONEY0
ws_ar.freeze_panes = f"A{AR_START}"
ws_ar.auto_filter.ref = f"B3:F{AR_END}"


# ==========================================================================
#              (10) بيانات نموذجية مبدئية (يستبدلها سكربت النقل)
# ==========================================================================
from datetime import datetime

# تعبئة قوائم الإعدادات بعناصر نموذجية
seed_lists = {
    "E": ["النقل الوطنية", "الشرق للتجارة", "دلتا للاستيراد", "المصرية للمعدات", "النيل للشحن"],  # الشركات
    "G": ["محمد علي", "أحمد سالم", "سيد رجب", "خالد فهمي"],                                      # السائقين
    "H": ["ورشة النصر", "مالك خاص", "شركة الطرق"],                                               # أصحاب السيارات
    "I": ["مصنع العبور", "شركة الأهرام", "التجارية الحديثة", "عميل الإسكندرية"],                 # العملاء
}
for col, vals in seed_lists.items():
    for i, v in enumerate(vals):
        ws_set[f"{col}{4+i}"] = v

def put(ws, row, colletter, val):
    ws[f"{colletter}{row}"] = val

# البوالص — 4 صفوف نموذجية
baw_samples = [
    (datetime(2026,1,12), "PL-1001", "المتوكل", "مستورد", "النقل الوطنية", "الإسكندرية", "CONT-4451", "ق ن ص 123", "محمد علي", "ورشة النصر",
     18000, 6000, 500, 300, 12000,   1500, 800, 300, 400, 200, 150, 100, 0, 250, 350, 200, 2500),
    (datetime(2026,2,20), "PL-1002", "المتوكل", "ترانزيت", "الشرق للتجارة", "العين السخنة", "CONT-4452", "ق ن ص 456", "أحمد سالم", "مالك خاص",
     22000, 8000, 0, 400, 15000,   1800, 900, 400, 500, 250, 200, 0, 300, 300, 500, 0, 3200),
    (datetime(2026,3,8), "PL-1003", "المتوكل", "بضائع عامة", "دلتا للاستيراد", "دمياط", "CONT-4453", "ق ن ص 789", "سيد رجب", "شركة الطرق",
     16000, 5000, 300, 0, 9000,   1200, 700, 0, 300, 180, 120, 90, 0, 200, 280, 150, 2100),
    (datetime(2026,5,15), "PL-1004", "المتوكل", "معدات", "المصرية للمعدات", "6 أكتوبر", "CONT-4454", "ق ن ص 321", "خالد فهمي", "ورشة النصر",
     25000, 9000, 600, 500, 18000,   2000, 1000, 500, 600, 300, 250, 200, 400, 350, 600, 300, 3800),
]
# ترتيب الأعمدة النموذجية: تاريخ، بوليصة، صاحب بيان، تصنيف، شركة، جهة، حاوية، رقم سيارة، سائق، صاحب سيارة،
#                          نولون، عهدة، تحويلة، بياتة تحصيل، دفعة، دور سواق، دور تباع، بياتة مصروف، مصروف، كرتات، تحيا مصر، ثبتان، ظ كاوتش، تصاريح، كاوتش، م أخرى، جاز
baw_map = ["التاريخ","رقم البوليصة","صاحب البيان","التصنيف","اسم الشركة","الجهة","رقم الحاوية","رقم السيارة","اسم السائق","صاحب السيارة",
           "نولون","عهدة","تحويلة","بياتة (تحصيل)","دفعة","دور سواق","دور تبّاع","بياتة (مصروف)","مصروف","كرتات","تحيا مصر","ثبتان","ظ كاوتش","تصاريح","كاوتش","م أخرى","جاز"]
for i, row in enumerate(baw_samples):
    rr = DATA_START + i
    for key, val in zip(baw_map, row):
        put(ws_baw, rr, Lb[key], val)

# التخليص — 3 صفوف
tak_samples = [
    (datetime(2026,1,20), "شهادة منشأ", "C-2001", 3500, 1200, 3000),
    (datetime(2026,2,11), "شهادة صحية", "C-2002", 2800, 900, 2000),
    (datetime(2026,4,5), "شهادة مطابقة", "C-2003", 4200, 1500, 4200),
]
tak_map = ["التاريخ","اسم الشهادة","رقم الشهادة","الأتعاب","المصروفات","الدفعات"]
for i, row in enumerate(tak_samples):
    rr = DATA_START + i
    for key, val in zip(tak_map, row):
        put(ws_tak, rr, Lt[key], val)

# المولدات
egy_samples = [
    (datetime(2026,1,9), "ايجي كول مصر", "المنطقة الصناعية - برج العرب", "PRM-501", 4, 12000, 8000),
    (datetime(2026,3,17), "ايجي كول مصر", "مدينة نصر", "PRM-502", 2, 6500, 6500),
    (datetime(2026,6,2), "ايجي كول مصر", "6 أكتوبر", "PRM-503", 6, 18000, 10000),
]
shg_samples = [
    (datetime(2026,2,14), "مؤسسة الشجيع", "طريق مصر إسكندرية الصحراوي", "SH-301", 3, 9000, 5000),
    (datetime(2026,5,22), "مؤسسة الشجيع", "العاشر من رمضان", "SH-302", 5, 15000, 15000),
]
gen_map = ["التاريخ","اسم الشركة","العنوان","رقم الإذن","عدد المولدات","النولون","الدفعات"]
for ws_x, samples in ((ws_egy, egy_samples), (ws_shg, shg_samples)):
    Lx = Le
    for i, row in enumerate(samples):
        rr = DATA_START + i
        for key, val in zip(gen_map, row):
            put(ws_x, rr, Lx[key], val)

# شامل — 4 فواتير
sha_samples = [
    (datetime(2026,1,25), "K-9001", "INV-7001", "PRM-A1", "مصنع العبور", "النقل الوطنية", "الإسكندرية", 3, 45000, 2000, 1500, 3000, 20000, 12000, 800, 5000),
    (datetime(2026,2,28), "K-9002", "INV-7002", "PRM-A2", "شركة الأهرام", "الشرق للتجارة", "دمياط", 2, 30000, 1000, 800, 2000, 12000, 8000, 500, 3500),
    (datetime(2026,4,12), "K-9003", "INV-7003", "PRM-A3", "التجارية الحديثة", "دلتا للاستيراد", "العين السخنة", 4, 60000, 3000, 2000, 4000, 30000, 16000, 1200, 6500),
    (datetime(2026,6,18), "K-9004", "INV-7004", "PRM-A4", "عميل الإسكندرية", "النيل للشحن", "بورسعيد", 1, 18000, 500, 0, 1000, 18000, 5000, 300, 2000),
]
sha_map = ["تاريخ الفاتورة","رقم الكشف","رقم الفاتورة","رقم الإذن","اسم العميل","اسم الشركة","الجهة","عدد الحاويات",
           "النولون","بيانات أو عطالات","خصم أ ت ص","الإيصالات","الدفعات","نولون فعلي","كشف وضمان","مصروف"]
for i, row in enumerate(sha_samples):
    rr = DATA_START + i
    for key, val in zip(sha_map, row):
        put(ws_sha, rr, Ls[key], val)


# ==========================================================================
#                       ترتيب الصفحات والحفظ
# ==========================================================================
order = [DB, "البوالص", "التخليص", "مولدات ايجي كول", "مولدات الشجيع", "شامل",
         AR, RPT, RP, S, GD]
wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else 99)
wb.active = wb.sheetnames.index(DB)

# إخفاء صفحة بيانات التقارير (محرك حسابات)
wb[RP].sheet_state = "hidden"

# لون تبويبات الصفحات
tabcolors = {DB: NAVY, "البوالص": GREEN, "التخليص": TEAL, "مولدات ايجي كول": PURPLE,
             "مولدات الشجيع": PURPLE, "شامل": GOLD, AR: "FF14746F", RPT: RED, S: "FF666666", GD: "FF999999"}
for name, col in tabcolors.items():
    wb[name].sheet_properties.tabColor = col[2:]

OUT = "/Users/karimmokhtar/المتوكل/المتوكل - نظام إدارة النقل والتخليص.xlsx"
wb.save(OUT)
print("تم الحفظ:", OUT)
print("عدد الصفحات:", len(wb.sheetnames))
print("الصفحات:", wb.sheetnames)


